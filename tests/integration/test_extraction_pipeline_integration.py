"""Integration tests: wire DrawingPreprocessor -> BalloonDetector -> ExtractionOrchestrator ->
ToleranceNormalizer together exactly as app.py's extraction routes do, using Fake* Azure clients
(no network calls) so the whole pipeline is exercised end-to-end in-process. Excel generation is
no longer part of the pipeline itself (see pipeline.py's module docstring) -- the reconciliation ->
export test below exercises ExcelWriter directly, the same way app.py's export route does.
"""
from __future__ import annotations

import io
from datetime import datetime, timezone

from openpyxl import load_workbook

from src.ai_clients import FakeChatCompletionClient, FakeDocumentAnalysisClient
from src.balloon_detector import BalloonDetector
from src.excel_writer import ExcelWriter
from src.exceptions import DocumentIntelligenceError, IncompleteReconciliationError, SegregationOfDutiesError
from src.extraction_orchestrator import ExtractionOrchestrator, VisionGroundedExtractionStrategy
from src.models import JobRecord, JobStatus, ReviewAction
from src.pipeline import ExtractionPipeline, PipelineContext
from src.preprocessor import DrawingPreprocessor
from src.reconciliation import ReconciliationService
from src.reconciliation_store import InMemoryReconciliationStore
from src.title_block_extractor import TitleBlockExtractor
from src.title_block_vision_reader import TitleBlockVisionReader
from src.tolerance_normalizer import ToleranceNormalizer

import pytest


def _build_pipeline(di_client, chat_client) -> ExtractionPipeline:
    return ExtractionPipeline(
        preprocessor=DrawingPreprocessor(target_dpi=150, min_dpi=72, max_pages=10),
        balloon_detector=BalloonDetector(di_client),
        orchestrator=ExtractionOrchestrator(VisionGroundedExtractionStrategy(chat_client)),
        normalizer=ToleranceNormalizer(),
        title_block_extractor=TitleBlockExtractor(),
        title_block_vision_reader=TitleBlockVisionReader(chat_client),
    )


def _job() -> JobRecord:
    return JobRecord(job_id="job-int-1", file_name="dwg.pdf", status=JobStatus.PROCESSING, created_at=datetime.now(timezone.utc))


def test_end_to_end_happy_path_produces_matching_counts(
    make_pdf_bytes, sample_layout, sample_extraction_response
):
    di_client = FakeDocumentAnalysisClient(sample_layout)
    chat_client = FakeChatCompletionClient([sample_extraction_response])
    pipeline = _build_pipeline(di_client, chat_client)

    ctx = PipelineContext(
        job=_job(), file_bytes=make_pdf_bytes(page_count=1), content_type="application/pdf",
        drawing_number="DWG-10245", revision="C",
    )
    ctx = pipeline.run(ctx)
    result = ExtractionPipeline.to_result(ctx, export_url=None)  # extraction never produces an export_url -- see app.py

    assert result.balloon_count_detected == 2  # balloons 12 and 13 from the fixture layout
    assert result.balloon_count_extracted == 2
    assert result.balloon_count_mismatch is False
    assert {b.balloon_number for b in result.balloons} == {12, 13}


def test_title_block_is_read_from_the_drawing_when_not_pre_supplied(make_pdf_bytes, sample_extraction_response):
    """FR-02 end-to-end: PipelineContext.drawing_number/.revision used to stay null forever --
    nothing in the pipeline ever set them. This reproduces the standard SIZE/DWG.NO./REV title
    block corner (SolidWorks' and most mechanical CAD tools' default sheet format), with each cell
    OCR'd as its own line -- confirmed against a real drawing to be how Document Intelligence
    actually segments this layout, not just a guess -- and confirms the pipeline now reads it, with
    no drawing_number/revision pre-supplied on the context this time."""
    layout_with_title_block = {
        "lines": [
            {"content": "TITLE"},
            {"content": "Counterweight Wheel"},
            {"content": "SIZE"},
            {"content": "DWG. NO."},
            {"content": "REV"},
            {"content": "B"},
            {"content": "CW-2045"},
            {"content": "A"},
            {"content": "SCALE: 1:1  WEIGHT: 604.55  SHEET 1 OF 1"},
        ],
        "words": [],
    }
    di_client = FakeDocumentAnalysisClient(layout_with_title_block)
    chat_client = FakeChatCompletionClient([sample_extraction_response])
    pipeline = _build_pipeline(di_client, chat_client)

    ctx = PipelineContext(job=_job(), file_bytes=make_pdf_bytes(page_count=1), content_type="application/pdf")
    ctx = pipeline.run(ctx)

    assert ctx.drawing_number == "CW-2045"
    assert ctx.revision == "A"


def test_title_block_falls_back_to_vision_when_the_regex_heuristic_finds_nothing(
    make_pdf_bytes, sample_extraction_response
):
    """A title block layout the regex heuristic doesn't recognize at all (no SIZE/DWG.NO./REV
    corner, no inline "LABEL: value" text) -- the regex correctly comes up empty, and the pipeline
    falls back to asking the vision model to read the title block directly off the page image."""
    layout_with_no_recognizable_title_block = {
        "lines": [
            {"content": "Widget 001"},
            {"content": "U.O.S Lengths ±0.25 Angles ±5°"},
            {"content": "Scale 1:1"},
        ],
        "words": [],
    }
    di_client = FakeDocumentAnalysisClient(layout_with_no_recognizable_title_block)
    chat_client = FakeChatCompletionClient(
        canned_responses=[sample_extraction_response],
        canned_structured=[{"drawing_number": "ABT-W001", "revision": None}],
    )
    pipeline = _build_pipeline(di_client, chat_client)

    ctx = PipelineContext(job=_job(), file_bytes=make_pdf_bytes(page_count=1), content_type="application/pdf")
    ctx = pipeline.run(ctx)

    assert ctx.drawing_number == "ABT-W001"
    assert ctx.revision is None  # this drawing genuinely has no revision field -- not a failure


def test_model_missing_a_detected_balloon_is_surfaced_not_dropped(
    make_pdf_bytes, sample_layout
):
    """FR-10 end-to-end: Document Intelligence detects 2 candidates on the page, but the AI model
    only returns 1 of them. The missing balloon must still appear in the final result (as a
    confidence=0 placeholder) and the mismatch flag must be set -- this is the exact signal the
    POC exists to measure."""
    di_client = FakeDocumentAnalysisClient(sample_layout)  # detects balloons 12 and 13
    partial_response = {
        "balloons": [
            {"balloon_number": 12, "page": 1, "nominal_value": 25.4, "unit": "mm", "confidence": 0.9}
        ]
    }
    chat_client = FakeChatCompletionClient([partial_response])
    pipeline = _build_pipeline(di_client, chat_client)

    ctx = PipelineContext(job=_job(), file_bytes=make_pdf_bytes(page_count=1), content_type="application/pdf")
    ctx = pipeline.run(ctx)
    result = ExtractionPipeline.to_result(ctx, export_url=None)

    assert result.balloon_count_detected == 2
    assert result.balloon_count_extracted == 1
    assert result.balloon_count_mismatch is True

    missing = next(b for b in result.balloons if b.balloon_number == 13)
    assert missing.confidence == 0.0
    assert missing.extraction_error is not None


def test_document_intelligence_outage_propagates_as_extraction_service_error(make_pdf_bytes):
    di_client = FakeDocumentAnalysisClient(RuntimeError("503 from upstream"))
    chat_client = FakeChatCompletionClient([])  # never reached
    pipeline = _build_pipeline(di_client, chat_client)

    ctx = PipelineContext(job=_job(), file_bytes=make_pdf_bytes(page_count=1), content_type="application/pdf")

    with pytest.raises(DocumentIntelligenceError):
        pipeline.run(ctx)


def test_multi_page_drawing_aggregates_balloons_across_pages(make_pdf_bytes):
    layout_page = {"lines": [], "words": [{"content": "1", "confidence": 0.9}]}
    di_client = FakeDocumentAnalysisClient(layout_page)
    responses = [
        {"balloons": [{"balloon_number": 1, "page": 1, "nominal_value": 5.0, "confidence": 0.9}]},
        {"balloons": [{"balloon_number": 1, "page": 2, "nominal_value": 7.0, "confidence": 0.9}]},
    ]
    chat_client = FakeChatCompletionClient(responses)
    pipeline = _build_pipeline(di_client, chat_client)

    ctx = PipelineContext(job=_job(), file_bytes=make_pdf_bytes(page_count=2), content_type="application/pdf")
    ctx = pipeline.run(ctx)

    assert len(ctx.balloons) == 2
    assert {b.page for b in ctx.balloons} == {1, 2}


def test_extraction_through_reconciliation_to_export_end_to_end(make_pdf_bytes, sample_layout, sample_extraction_response):
    """The full loop requirements.md's "quality check/reconciliation to ensure 100% data
    accuracy" item describes: extract -> nothing exportable yet -> every balloon reviewed ->
    signed off -> export now reflects the human-verified values, not the raw AI output."""
    di_client = FakeDocumentAnalysisClient(sample_layout)
    chat_client = FakeChatCompletionClient([sample_extraction_response])
    pipeline = _build_pipeline(di_client, chat_client)
    reconciliation = ReconciliationService(InMemoryReconciliationStore())

    ctx = PipelineContext(
        job=_job(), file_bytes=make_pdf_bytes(page_count=1), content_type="application/pdf",
        drawing_number="DWG-10245", revision="C",
    )
    ctx = pipeline.run(ctx)  # balloons 12 (nominal 25.4) and 13 (a GD&T-only balloon), per the fixture

    record = reconciliation.start(
        "job-int-1", ctx.drawing_number, ctx.revision, ctx.balloons, submitted_by="alice", template_id="generic-flat"
    )
    assert all(b.status.value == "pending" for b in record.balloons)
    assert record.template_id == "generic-flat"  # remembered from upload, used as export's default

    # export is blocked before any review happens at all
    with pytest.raises(IncompleteReconciliationError):
        reconciliation.get_reconciled_balloons("job-int-1")

    # the analyst who submitted it cannot also be the reviewer
    with pytest.raises(SegregationOfDutiesError):
        reconciliation.review_balloon("job-int-1", page=1, balloon_number=12, reviewer_id="alice", action=ReviewAction.CONFIRM)

    reconciliation.review_balloon("job-int-1", page=1, balloon_number=12, reviewer_id="bob", action=ReviewAction.CONFIRM)
    # reviewer catches an AI transcription error on the second balloon and corrects it
    corrected = next(b for b in ctx.balloons if b.balloon_number == 13).model_copy(update={"nominal_value": 12.7})
    reviewed = reconciliation.review_balloon(
        "job-int-1", page=1, balloon_number=13, reviewer_id="bob", action=ReviewAction.CORRECT, corrected=corrected
    )
    assert reviewed.discrepancy is True

    status = reconciliation.get_status("job-int-1")
    assert status.percent_complete == 100.0
    assert status.ready_for_signoff is True

    # signing off before 100% would 409; here it succeeds since both balloons are now reconciled
    reconciliation.sign_off("job-int-1", "bob")

    finalized = reconciliation.get_reconciled_balloons("job-int-1")
    extracted = reconciliation.get_extracted_balloons("job-int-1")
    record = reconciliation.get_record("job-int-1")
    excel_bytes = ExcelWriter().write(
        drawing_number=ctx.drawing_number,
        revision=ctx.revision,
        reconciled_balloons=finalized,
        extracted_balloons=extracted,
        template_id=record.template_id,
    )

    corrected_row = next(b for b in finalized if b.balloon_number == 13)
    assert corrected_row.nominal_value == 12.7  # the reviewer's correction, not the AI's original value
    assert corrected_row.confidence == 1.0  # human-verified
    assert all(b.confidence == 1.0 for b in finalized)

    wb = load_workbook(io.BytesIO(excel_bytes))
    assert wb.sheetnames == ["Reconciled", "Extracted"]
    # generic-flat has no title block for either tab
    assert wb["Reconciled"].max_row == 3  # header + 2 reconciled balloon rows
    assert wb["Extracted"].max_row == 3  # header + 2 raw-extracted balloon rows

    # the Extracted tab shows the AI's original (unreviewed) value for balloon 13 -- it never had a
    # nominal_value at all, just a GD&T frame (see sample_extraction_response.json); the Reconciled
    # tab shows the reviewer's correction instead.
    nominal_col = 5  # balloon_number, drawing_number, revision, page, nominal_value (see excel_templates.GENERIC_FLAT)
    extracted_row = next(r for r in wb["Extracted"].iter_rows(min_row=2) if r[0].value == 13)
    reconciled_row = next(r for r in wb["Reconciled"].iter_rows(min_row=2) if r[0].value == 13)
    assert extracted_row[nominal_col - 1].value is None
    assert reconciled_row[nominal_col - 1].value == 12.7
