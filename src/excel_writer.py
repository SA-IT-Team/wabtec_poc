"""ExcelWriter: renders the export workbook for whichever named template was chosen (FR-17,
src/excel_templates.py) -- resolving the template and mapping balloon data onto its columns is the
registry's job, this class just lays cells out on worksheets from what the registry hands back.

The workbook always carries a "Reconciled" tab (the human-verified export -- FR-19) and, when raw
extraction data is supplied, a second "Extracted" tab with the AI's original output alongside it,
so a reviewer or auditor can see what the model actually produced versus what a human confirmed or
corrected without cross-referencing a separate file (FR-27 traceability). Both tabs use the same
template's column layout -- the template governs *how* each tab is laid out, not which data lands
on which tab.
"""
from __future__ import annotations

import io
from datetime import datetime, timezone

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from src.excel_templates import ExcelTemplate, get_template
from src.models import ExtractedBalloon

LOW_CONFIDENCE_THRESHOLD = 0.7
HEADER_FILL = PatternFill(start_color="1C2431", end_color="1C2431", fill_type="solid")
LOW_CONFIDENCE_FILL = PatternFill(start_color="F3E7C9", end_color="F3E7C9", fill_type="solid")
TITLE_LABEL_FONT = Font(bold=True)

RECONCILED_SHEET_NAME = "Reconciled"
EXTRACTED_SHEET_NAME = "Extracted"


class ExcelWriter:
    def write(
        self,
        *,
        drawing_number: str | None,
        revision: str | None,
        reconciled_balloons: list[ExtractedBalloon],
        extracted_balloons: list[ExtractedBalloon] | None = None,
        template_id: str | None = None,
    ) -> bytes:
        """`reconciled_balloons` (human-verified, e.g. ReconciliationService.get_reconciled_balloons)
        always becomes the active "Reconciled" tab. Pass `extracted_balloons` (the raw AI output,
        e.g. ReconciliationService.get_extracted_balloons) to also get a second "Extracted" tab for
        comparison -- omit it for a single-tab workbook (e.g. a quick preview before sign-off).

        Raises UnknownTemplateError (-> HTTP 400 in app.py) if template_id doesn't match a
        registered template -- see src/excel_templates.py."""
        template = get_template(template_id)

        wb = Workbook()
        reconciled_ws = wb.active
        reconciled_ws.title = RECONCILED_SHEET_NAME
        self._write_sheet(reconciled_ws, template, reconciled_balloons, drawing_number, revision)

        if extracted_balloons is not None:
            extracted_ws = wb.create_sheet(EXTRACTED_SHEET_NAME)
            self._write_sheet(extracted_ws, template, extracted_balloons, drawing_number, revision)

        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    @classmethod
    def _write_sheet(
        cls,
        ws: Worksheet,
        template: ExcelTemplate,
        balloons: list[ExtractedBalloon],
        drawing_number: str | None,
        revision: str | None,
    ) -> None:
        header_row = cls._write_title_block(ws, template, drawing_number, revision) if template.title_block else 1
        cls._write_header(ws, template, header_row)
        cls._write_rows(ws, template, balloons, drawing_number, revision, header_row)

        for col_idx in range(1, len(template.columns) + 1):
            ws.column_dimensions[cls._column_letter(col_idx)].width = 20 if template.title_block else 16

    @staticmethod
    def _write_title_block(ws, template: ExcelTemplate, drawing_number: str | None, revision: str | None) -> int:
        rows = [
            ("Drawing / Part No.", drawing_number or "—"),
            ("Revision", revision or "—"),
            ("Report Generated", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")),
        ]
        for row_idx, (label, value) in enumerate(rows, start=1):
            ws.cell(row=row_idx, column=1, value=label).font = TITLE_LABEL_FONT
            ws.cell(row=row_idx, column=2, value=value)
        return len(rows) + 2  # one blank row between the title block and the table header

    @staticmethod
    def _write_header(ws, template: ExcelTemplate, header_row: int) -> None:
        for col_idx, (_, header) in enumerate(template.columns, start=1):
            cell = ws.cell(row=header_row, column=col_idx, value=header)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = HEADER_FILL

    @staticmethod
    def _write_rows(
        ws,
        template: ExcelTemplate,
        balloons: list[ExtractedBalloon],
        drawing_number: str | None,
        revision: str | None,
        header_row: int,
    ) -> None:
        for offset, balloon in enumerate(sorted(balloons, key=lambda b: b.balloon_number), start=1):
            row_idx = header_row + offset
            row_values = template.row_values(balloon, drawing_number, revision)
            for col_idx, (key, _) in enumerate(template.columns, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=row_values.get(key))
                if balloon.confidence < LOW_CONFIDENCE_THRESHOLD:
                    cell.fill = LOW_CONFIDENCE_FILL

    @staticmethod
    def _column_letter(col_idx: int) -> str:
        letters = ""
        while col_idx > 0:
            col_idx, remainder = divmod(col_idx - 1, 26)
            letters = chr(65 + remainder) + letters
        return letters
